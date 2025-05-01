import os
import logging
from typing import List, Dict, Any, Optional
import json
from pathlib import Path
import docx
from tqdm import tqdm
import torch

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_core.documents import Document
from dotenv import load_dotenv

load_dotenv()

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

class DocumentEmbedder:
    
    def __init__(self, 
                 embedding_model_name: str = "sentence-transformers/all-mpnet-base-v2",
                 chunk_size: int = 1000,
                 chunk_overlap: int = 200,
                 vector_store_path: str = "artifacts/vector_store",
                 use_gpu: bool = None):
        
        self.embedding_model_name = embedding_model_name
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.vector_store_path = vector_store_path
        
        if use_gpu is None:
            use_gpu = torch.cuda.is_available()
        
        device = "cuda" if use_gpu else "cpu"
        logger.info(f"Initializing embedder with model: {embedding_model_name} on device: {device}")
        
        self.embeddings = HuggingFaceEmbeddings(
            model_name=embedding_model_name,
            model_kwargs={"device": device}
        )
        
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=len
        )
        
        if use_gpu:
            logger.info(f"Using GPU acceleration: {torch.cuda.get_device_name(0)}")
            logger.info(f"CUDA Version: {torch.version.cuda}")
        else:
            logger.info("Running on CPU. GPU not available or not enabled.")
        
    def read_docx(self, file_path: str) -> str:
        logger.info(f"Reading DOCX file: {file_path}")
        try:
            doc = docx.Document(file_path)
            full_text = []
            
            for para in doc.paragraphs:
                full_text.append(para.text)
                
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        full_text.append(cell.text)
            
            return "\n".join(full_text)
        except Exception as e:
            logger.error(f"Error reading DOCX file: {e}")
            raise
    
    def process_document(self, file_path: str, metadata: Optional[Dict[str, Any]] = None) -> List[Document]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Document not found: {file_path}")
            
        logger.info(f"Processing document: {file_path}")
        
        text_content = self.read_docx(file_path)
        
        doc_metadata = metadata or {}
        doc_metadata["source"] = file_path
        doc_metadata["filename"] = os.path.basename(file_path)
        
        document = Document(page_content=text_content, metadata=doc_metadata)
        
        logger.info(f"Splitting document into chunks (size: {self.chunk_size}, overlap: {self.chunk_overlap})")
        chunks = self.text_splitter.split_documents([document])
        
        logger.info(f"Created {len(chunks)} chunks from document")
        return chunks
    
    def create_embeddings(self, documents: List[Document]) -> FAISS:
        logger.info(f"Creating embeddings for {len(documents)} documents")
        vectorstore = FAISS.from_documents(documents, self.embeddings)
        return vectorstore
    
    def save_vector_store(self, vectorstore: FAISS, store_name: str = "docx_vectors") -> str:
        os.makedirs(self.vector_store_path, exist_ok=True)
        
        save_path = os.path.join(self.vector_store_path, store_name)
        vectorstore.save_local(save_path)
        logger.info(f"Vector store saved to {save_path}")
        
        return save_path
    
    def load_vector_store(self, store_name: str = "docx_vectors") -> FAISS:
        load_path = os.path.join(self.vector_store_path, store_name)
        if not os.path.exists(load_path):
            raise FileNotFoundError(f"Vector store not found at {load_path}")
            
        logger.info(f"Loading vector store from {load_path}")
        return FAISS.load_local(load_path, self.embeddings)
    
    def embed_docx_file(self, file_path: str, store_name: str = None) -> str:
        if store_name is None:
            store_name = os.path.splitext(os.path.basename(file_path))[0] + "_vectors"
        
        chunks = self.process_document(file_path)
        
        vectorstore = self.create_embeddings(chunks)
        return self.save_vector_store(vectorstore, store_name)
    
    def embed_multiple_docx_files(self, file_paths: List[str], store_name: str = "combined_vectors") -> str:
        all_chunks = []
        
        for file_path in tqdm(file_paths, desc="Processing documents"):
            chunks = self.process_document(file_path)
            all_chunks.extend(chunks)
        
        logger.info(f"Created total of {len(all_chunks)} chunks from {len(file_paths)} documents")
        
        vectorstore = self.create_embeddings(all_chunks)
        return self.save_vector_store(vectorstore, store_name)


def main():
    gpu_available = torch.cuda.is_available()
    if gpu_available:
        logger.info(f"GPU detected: {torch.cuda.get_device_name(0)}")
        logger.info(f"Running embedding with GPU acceleration")
    else:
        logger.warning("No GPU detected. Running on CPU which may be slower.")
    
    embedder = DocumentEmbedder(
        chunk_size=1000,
        chunk_overlap=200,
        vector_store_path="artifacts/vector_store",
        use_gpu=True
    )
    
    docx_path = "combined_rag_data.docx"
    
    if os.path.exists(docx_path):
        store_path = embedder.embed_docx_file(docx_path, "combined_rag_vectors")
        logger.info(f"Embeddings created and saved to {store_path}")
    else:
        logger.error(f"File not found: {docx_path}")


if __name__ == "__main__":
    main()